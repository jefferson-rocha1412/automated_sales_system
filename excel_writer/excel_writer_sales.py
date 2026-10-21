import pandas as pd
import os
from sales.sales_data import get_excel_data
from config.config import Config
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from utils.logger import logger_setup

def sales_report():
    # CREATE THE LOGGER INFO
    logger = logger_setup()
    
    # BORDER SETTINGS
    thick_border = Border(
        left=Side(style='thick', color='000000'),
        right=Side(style='thick', color='000000'),
        top=Side(style='thick', color='000000'),
        bottom=Side(style='thick', color='000000')
    )

    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )

    sales_data, sales_by_person = get_excel_data()

    if sales_data is None or sales_by_person is None:
        return None

    # CREATE THE SAVE PATH
    save_path = os.path.join(Config.OUTPUT_FOLDER, Config.SALES_REPORT_WORKBOOK_FILENAME)

    try:
        with pd.ExcelWriter(save_path) as writer:
            sales_data.to_excel(
                writer,
                sheet_name=Config.SALES_REPORT_SUMMARY_SHEET,
                index=False
                )
            sales_by_person.to_excel(
                writer,
                sheet_name=Config.SALES_REPORT_SALESBYPERSON_SHEET,
                index=False
            )

            #-- FORMATTING THE SUMMARY SHEET --#
            worksheet_summary = writer.sheets[Config.SALES_REPORT_SUMMARY_SHEET]

            # FORMAT THE HEADERS
            for cell in worksheet_summary[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.alignment = Alignment(horizontal='center')
                cell.fill = PatternFill(fill_type='solid', fgColor='1F4E78')
                cell.border = thick_border

                if cell.value == 'Unit Price':
                    unitprice_column_letter = get_column_letter(cell.column)

                if cell.value == 'Sales':
                    sales_column_letter = get_column_letter(cell.column)

            # FORMAT THE PRICE COLUMN
            for cell in worksheet_summary[unitprice_column_letter][1:]:
                cell.number_format = "#,##0.00"

            # FORMAT THE SALES COLUMN
            for cell in worksheet_summary[sales_column_letter][1:]:
                cell.number_format = "#,##0.00"

            # FORMAT THE WIDTH
            for column_number, column in enumerate(worksheet_summary.columns, start=1):
                max_length = 0

                # GET THE COLUMN LETTER
                column_letter = get_column_letter(column_number)

                for cell in column:
                    if cell.value is not None:
                        if cell.row == 1:
                            length = len(str(cell.value))
                        elif column_letter in [unitprice_column_letter, sales_column_letter]:
                            length = len(f"{cell.value:,.2f}")
                        else:
                            length = len(str(cell.value))
                            cell.alignment = Alignment(horizontal='center', wrap_text=True)

                        max_length = max(max_length, length)

                worksheet_summary.column_dimensions[column_letter].width = max_length + 6

            # FREEZE PANES
            worksheet_summary.freeze_panes = 'A2'

            # AUTO FILTERING
            worksheet_summary.auto_filter.ref = worksheet_summary.dimensions

            #-- FORMATTING THE SALES BY PERSON SHEET --#
            worksheet_salesbyperson = writer.sheets[Config.SALES_REPORT_SALESBYPERSON_SHEET]

            # FORMAT THE HEADERS
            for cell in worksheet_salesbyperson[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.alignment = Alignment(horizontal='center')
                cell.fill = PatternFill(fill_type='solid', fgColor='FF8C00')
                cell.border = thick_border

                if cell.value == 'Sales':
                    sales2_column_letter = get_column_letter(cell.column)

            # FORMAT THE SALES COLUMN
            for cell in worksheet_salesbyperson[sales2_column_letter][1:]:
                cell.number_format = "#,##0.00"

            # FORMAT THE WIDTH
            for column_number, column in enumerate(worksheet_salesbyperson.columns, start=1):
                max_length = 0

                # GET THE COLUMN LETTER
                column_letter = get_column_letter(column_number)

                for cell in column:
                    cell.border = thin_border

                    if cell.value is not None:
                        if cell.row == 1:
                            length = len(str(cell.value))
                        elif column_letter == sales2_column_letter:
                            length = len(f"{cell.value:,.2f}")
                        else:
                            length = len(str(cell.value))
                        
                        max_length = max(max_length, length)

                worksheet_salesbyperson.column_dimensions[column_letter].width = max_length + 6
            
    except PermissionError:
        logger.info(f'Unable to save {Config.SALES_REPORT_WORKBOOK_FILENAME}. Please close the file and try again.')

if __name__ == "__main__":
    sales_report()